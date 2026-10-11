"""Dashboard, laporan (hanya transaksi AKTIF), ekspor PDF/Excel, dan audit log."""
import io
import sqlite3
from datetime import date, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Response

from . import config
from .db import get_db, now
from .deps import current_user, require_owner
from .routes_tx import range_bounds, rp, tx_summary

router = APIRouter(prefix="/api")
ACTIVE = "t.status='COMPLETED'"


def _sum(conn, start, end, extra="", args=()):
    r = conn.execute(
        "SELECT COUNT(*) AS n, COALESCE(SUM(t.total),0) AS sales, COALESCE(SUM(t.payment),0) AS paid, "
        f"COALESCE(SUM(t.change),0) AS chg FROM transactions t WHERE {ACTIVE} "
        f"AND t.transaction_date>=? AND t.transaction_date<? {extra}", (start, end, *args)).fetchone()
    return {"count": r["n"], "sales": r["sales"], "paid": r["paid"], "change": r["chg"]}


def _top_items(conn, start, end, limit=10):
    rows = conn.execute(
        "SELECT i.product_name_snapshot AS name, i.category_name_snapshot AS category, "
        "SUM(i.quantity) AS quantity, SUM(i.subtotal) AS revenue FROM transaction_items i "
        f"JOIN transactions t ON t.id=i.transaction_id WHERE {ACTIVE} AND t.transaction_date>=? "
        "AND t.transaction_date<? GROUP BY i.product_name_snapshot, i.category_name_snapshot "
        "ORDER BY quantity DESC, revenue DESC LIMIT ?", (start, end, limit)).fetchall()
    return [{"name": r["name"], "category": r["category"], "quantity": r["quantity"],
             "revenue": r["revenue"]} for r in rows]


def _daily(conn, d_from: date, d_to: date):
    start, end = range_bounds(d_from, d_to)
    rows = conn.execute(
        "SELECT substr(t.transaction_date,1,10) AS d, COUNT(*) AS n, SUM(t.total) AS s "
        f"FROM transactions t WHERE {ACTIVE} AND t.transaction_date>=? AND t.transaction_date<? "
        "GROUP BY d", (start, end)).fetchall()
    by = {r["d"]: (r["n"], r["s"]) for r in rows}
    out, d = [], d_from
    while d <= d_to:
        n, s = by.get(d.isoformat(), (0, 0))
        out.append({"date": d.isoformat(), "count": n, "total": s})
        d += timedelta(days=1)
    return out


# ---------------- Dashboard ----------------
@router.get("/dashboard/owner")
def dashboard_owner(owner=Depends(require_owner), conn: sqlite3.Connection = Depends(get_db)):
    today = now().date()
    m_start = today.replace(day=1)
    m_next = (m_start + timedelta(days=32)).replace(day=1)
    t0, t1 = range_bounds(today, today)
    m0, m1 = range_bounds(m_start, m_next - timedelta(days=1))
    day, month = _sum(conn, t0, t1), _sum(conn, m0, m1)
    month_all = conn.execute("SELECT COUNT(*) FROM transactions WHERE transaction_date>=? AND transaction_date<?",
                             (m0, m1)).fetchone()[0]
    recent = conn.execute(
        "SELECT t.*, (SELECT COUNT(*) FROM transaction_items i WHERE i.transaction_id=t.id) AS item_count "
        "FROM transactions t WHERE t.status='COMPLETED' ORDER BY t.transaction_date DESC, t.id DESC LIMIT 8").fetchall()
    deleted = conn.execute(
        "SELECT t.* FROM transactions t WHERE t.status='DELETED' ORDER BY t.deleted_at DESC LIMIT 5").fetchall()
    return {"today_sales": day["sales"], "today_count": day["count"],
            "month_sales": month["sales"], "month_active_count": month["count"],
            "month_transaction_count": month_all, "month_deleted_count": month_all - month["count"],
            "top_items": _top_items(conn, m0, m1, 5),
            "chart": _daily(conn, today - timedelta(days=13), today),
            "recent": [tx_summary(r) for r in recent],
            "deleted": [tx_summary(r) for r in deleted]}


@router.get("/dashboard/cashier")
def dashboard_cashier(user=Depends(current_user), conn: sqlite3.Connection = Depends(get_db)):
    t0, t1 = range_bounds(now().date(), now().date())
    s = _sum(conn, t0, t1, "AND t.cashier_id=?", (user["id"],))
    return {"today_count": s["count"], "today_total": s["sales"]}


# ---------------- Laporan ----------------
def _period(d_from: Optional[date], d_to: Optional[date]):
    today = now().date()
    d_from, d_to = d_from or today, d_to or d_from or today
    if d_from > d_to:
        raise HTTPException(400, "Tanggal awal tidak boleh setelah tanggal akhir.")
    if (d_to - d_from).days > 366:
        raise HTTPException(400, "Rentang laporan maksimal 1 tahun.")
    return d_from, d_to


def build_summary(conn, d_from: date, d_to: date) -> dict:
    start, end = range_bounds(d_from, d_to)
    s = _sum(conn, start, end)
    deleted = conn.execute("SELECT COUNT(*) FROM transactions WHERE status='DELETED' AND "
                           "transaction_date>=? AND transaction_date<?", (start, end)).fetchone()[0]
    by_method = conn.execute(
        f"SELECT t.payment_method AS m, COUNT(*) AS n, SUM(t.total) AS s FROM transactions t WHERE {ACTIVE} "
        "AND t.transaction_date>=? AND t.transaction_date<? GROUP BY t.payment_method", (start, end)).fetchall()
    by_cat = conn.execute(
        "SELECT i.category_name_snapshot AS c, SUM(i.subtotal) AS s FROM transaction_items i JOIN transactions t "
        f"ON t.id=i.transaction_id WHERE {ACTIVE} AND t.transaction_date>=? AND t.transaction_date<? "
        "GROUP BY i.category_name_snapshot ORDER BY s DESC", (start, end)).fetchall()
    return {"from": d_from.isoformat(), "to": d_to.isoformat(),
            "total_sales": s["sales"], "transaction_count": s["count"], "active_count": s["count"],
            "deleted_count": deleted, "total_payment": s["paid"], "total_change": s["change"],
            "average": (s["sales"] // s["count"]) if s["count"] else 0,
            "top_items": _top_items(conn, start, end, 10),
            "daily": _daily(conn, d_from, d_to),
            "by_category": [{"category": r["c"], "total": r["s"]} for r in by_cat],
            "by_payment_method": [{"method": r["m"], "count": r["n"], "total": r["s"]} for r in by_method]}


@router.get("/reports/summary")
def report_summary(from_date: Optional[date] = None, to_date: Optional[date] = None,
                   owner=Depends(require_owner), conn: sqlite3.Connection = Depends(get_db)):
    d_from, d_to = _period(from_date, to_date)
    return build_summary(conn, d_from, d_to)


def _tx_rows(conn, d_from, d_to):
    start, end = range_bounds(d_from, d_to)
    return conn.execute(f"SELECT t.* FROM transactions t WHERE {ACTIVE} AND t.transaction_date>=? "
                        "AND t.transaction_date<? ORDER BY t.transaction_date", (start, end)).fetchall()


@router.get("/reports/export.xlsx")
def export_xlsx(from_date: Optional[date] = None, to_date: Optional[date] = None,
                owner=Depends(require_owner), conn: sqlite3.Connection = Depends(get_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    d_from, d_to = _period(from_date, to_date)
    s = build_summary(conn, d_from, d_to)
    wb = Workbook()
    ws = wb.active; ws.title = "Ringkasan"
    head = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="0B1957")
    ws["A1"] = config.CLINIC_NAME; ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Laporan penjualan {d_from.isoformat()} s/d {d_to.isoformat()} (hanya transaksi aktif)"
    rows = [("Total penjualan", s["total_sales"]), ("Jumlah transaksi aktif", s["transaction_count"]),
            ("Total pembayaran diterima", s["total_payment"]), ("Total kembalian", s["total_change"]),
            ("Rata-rata per transaksi", s["average"])]
    for i, (k, v) in enumerate(rows, start=4):
        ws.cell(i, 1, k); ws.cell(i, 2, v)
        if "transaksi" not in k.lower():
            ws.cell(i, 2).number_format = '"Rp"#,##0'
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 20
    t = wb.create_sheet("Transaksi")
    t.append(["No. Kwitansi", "Tanggal", "Pasien", "Kasir", "Cara Bayar", "Total", "Bayar", "Kembali"])
    for r in _tx_rows(conn, d_from, d_to):
        t.append([r["transaction_number"], r["transaction_date"][:16].replace("T", " "), r["patient_name"],
                  r["cashier_name"], r["payment_method"], r["total"], r["payment"], r["change"]])
    i = wb.create_sheet("Item Terjual")
    i.append(["Item", "Kategori", "Jumlah", "Pendapatan"])
    for it in _top_items(conn, *range_bounds(d_from, d_to), limit=1000):
        i.append([it["name"], it["category"], it["quantity"], it["revenue"]])
    for sh in (t, i):
        for c in sh[1]:
            c.font = head; c.fill = fill; c.alignment = Alignment(horizontal="center")
        for col, w in zip("ABCDEFGH", (16, 18, 28, 22, 12, 14, 14, 14)):
            sh.column_dimensions[col].width = w
        for row in sh.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, int) and c.column_letter in "FGH" and sh.title == "Transaksi":
                    c.number_format = '"Rp"#,##0'
        if sh.title == "Item Terjual":
            for c in sh["D"][1:]:
                c.number_format = '"Rp"#,##0'
    buf = io.BytesIO(); wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="Laporan_{d_from}_{d_to}.xlsx"'})


@router.get("/reports/export.pdf")
def export_pdf(from_date: Optional[date] = None, to_date: Optional[date] = None,
               owner=Depends(require_owner), conn: sqlite3.Connection = Depends(get_db)):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    d_from, d_to = _period(from_date, to_date)
    s = build_summary(conn, d_from, d_to)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm,
                            title=f"Laporan {d_from} - {d_to}")
    st = getSampleStyleSheet()
    navy = colors.HexColor("#0B1957")

    def table(data, widths, align_right=()):
        t = Table(data, colWidths=widths, repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), navy), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                 ("FONTSIZE", (0, 0), (-1, -1), 8), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#9ECCFA")),
                 ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8F3EA")])]
        style += [("ALIGN", (c, 0), (c, -1), "RIGHT") for c in align_right]
        t.setStyle(TableStyle(style)); return t
    story = [Paragraph(config.CLINIC_NAME, st["Title"]),
             Paragraph(f"Laporan penjualan {d_from.strftime('%d/%m/%Y')} - {d_to.strftime('%d/%m/%Y')} "
                       "(transaksi aktif saja)", st["Normal"]), Spacer(1, 6 * mm)]
    story.append(table([["Ringkasan", ""], ["Total penjualan", rp(s["total_sales"])],
                        ["Jumlah transaksi aktif", str(s["transaction_count"])],
                        ["Total pembayaran diterima", rp(s["total_payment"])],
                        ["Rata-rata per transaksi", rp(s["average"])]], [70 * mm, 50 * mm], (1,)))
    story += [Spacer(1, 6 * mm), Paragraph("Item terlaris", st["Heading3"])]
    story.append(table([["Item", "Kategori", "Jumlah", "Pendapatan"]] +
                       [[x["name"][:42], x["category"], str(x["quantity"]), rp(x["revenue"])]
                        for x in s["top_items"]], [75 * mm, 30 * mm, 20 * mm, 40 * mm], (2, 3)))
    story += [Spacer(1, 6 * mm), Paragraph("Daftar transaksi", st["Heading3"])]
    story.append(table([["No.", "Tanggal", "Pasien", "Kasir", "Total"]] +
                       [[r["transaction_number"], r["transaction_date"][:16].replace("T", " "),
                         r["patient_name"][:28], r["cashier_name"][:20], rp(r["total"])]
                        for r in _tx_rows(conn, d_from, d_to)], [22 * mm, 33 * mm, 55 * mm, 40 * mm, 30 * mm], (4,)))
    doc.build(story)
    return Response(buf.getvalue(), media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="Laporan_{d_from}_{d_to}.pdf"'})


# ---------------- Audit log ----------------
@router.get("/audit-logs")
def audit_logs(action: str = "", q: str = "", limit: int = 50, offset: int = 0,
               owner=Depends(require_owner), conn: sqlite3.Connection = Depends(get_db)):
    limit = max(1, min(limit, 200)); offset = max(0, offset)
    where, args = [], []
    if action:
        where.append("action=?"); args.append(action)
    if q.strip():
        where.append("(description LIKE ? OR username LIKE ?)"); args += [f"%{q.strip()}%"] * 2
    w = ("WHERE " + " AND ".join(where)) if where else ""
    total = conn.execute(f"SELECT COUNT(*) FROM audit_logs {w}", args).fetchone()[0]
    rows = conn.execute(f"SELECT * FROM audit_logs {w} ORDER BY id DESC LIMIT ? OFFSET ?", (*args, limit, offset)).fetchall()
    return {"total": total, "items": [{k: r[k] for k in r.keys()} for r in rows]}
