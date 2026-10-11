#!/usr/bin/env python3
"""Baca data-Alqoyyum-royal-clinic.xlsx -> seed/products.json + seed/import_report.json

Sumber utama : sheet HARGA (B = nama, C = harga modal/HPP, D = harga jual).
Tambahan     : item yang hanya muncul di sheet kwitansi (Kuretase, lahiran, dst.).
Nilai kosong/#N/A/0 pada harga modal TIDAK dipakai (disimpan null). Tidak ada crash.

Pakai: python tools/extract_excel.py [path.xlsx]
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from app.icons import pick_icon  # noqa: E402

DEFAULT_SRC = ROOT / "seed" / "data-Alqoyyum-royal-clinic.xlsx"

# nama, label di struk (sesuai Excel), urutan di struk, urutan tampil
CATEGORIES = [
    ("Obat", "Obat", 4, 1),
    ("Jasa", "Jasa", 1, 2),
    ("Tindakan", "Jasa", 1, 3),     # di Excel, tindakan dicetak di bawah label "Jasa"
    ("Lab", "Lab", 2, 4),
    ("BHP", "BHP", 5, 5),           # Bahan Habis Pakai (label "BHP" di sheet Kuretase)
    ("Alat", "Alat", 3, 6),
]
CATEGORY_ICON = {"Obat": "medicine/medicine", "Jasa": "service/service",
                 "Tindakan": "treatment/procedure", "Lab": "service/lab",
                 "BHP": "equipment/equipment", "Alat": "equipment/equipment"}

JASA = {"konsultasi dokter", "usg 2d", "usg 3d", "usg 4d", "transvaginal"}
TINDAKAN = {"papsmear", "inspekulo", "periksa dalam", "gv", "gv sedang", "gv kecil",
            "irigasi vagina", "insersi iud", "ekstraksi iud", "injeksi"}
OVERRIDE = {"nasal aspirator": "Alat", "nova t iud": "Alat", "ultra sonic gel": "BHP",
            "buku kia": "BHP", "betadine": "BHP"}
ALAT_BLOCK = {"torniquet onemed", "gunting verban sakamed", "mortar 100 mm sella",
              "ring pesarium no.7", "tempat korentang", "korentang", "pulse oximeter a2"}
AMPUL = {"ceftriaxone sodium 1 g inj", "oxytocin inj", "wi 20 ml (aqua pro injection)"}

# (nama di app, nama sumber di Excel, kategori, harga, harga_bebas, sheet sumber)
EXTRAS = [
    ("Dokter Kandungan", "Dokter Kandungan", "Jasa", 1500000, True, "Kuretase"),
    ("Dokter Anestesi & Obat Anestesi", "Dokter Anestesi & Obat Anestesi", "Jasa", 900000, True, "Kuretase"),
    ("Jasa Keperawatan", "Jasa Keperawatan", "Jasa", 400000, True, "Kuretase"),
    ("Administrasi", "Administrasi", "Jasa", 50000, False, "Kuretasi (2)"),
    ("One Day Care", "One Day Care", "Jasa", 500000, True, "Kuretasi (2)"),
    ("Ruangan", "Ruangan", "Jasa", 1300000, True, "lahiran"),
    ("Lain-lain", "Lain-lain", "Jasa", 0, True, "lahiran"),
    ("Tindakan Kuretasi", "Tindakan Kuretasi", "Tindakan", 2100000, True, "Kuretase"),
    ("Tindakan Infus", "Tindakan Infus", "Tindakan", 25000, False, "Kuretase"),
    ("Biaya Tindakan", "Biaya Tindakan", "Tindakan", 100000, True, "mabok.ngidam"),
    ("Jasa Tindakan (Paket Persalinan)", "Jasa Tindakan", "Tindakan", 4250000, True, "lahiran"),
    ("Glukosa Sewaktu", "Glukosa Sewaktu", "Lab", 57000, False, "Kuretase"),
    ("HBsAG", "HBsAG", "Lab", 277000, False, "Kuretase"),
    ("Hematologi Rutin", "Hematologi Rutin", "Lab", 176000, False, "Kuretase"),
    ("Infusan RL", "Infusan RL", "Obat", 35000, False, "mabok.ngidam"),
    ("Obat & Alkes (Paket)", "Obat & Alkes", "Obat", 0, True, "Kuretasi (2)"),
    ("Alat Kesehatan (Paket)", "Alat Kesehatan", "Alat", 0, True, "lahiran"),
    ("Kateter", "Kateter", "BHP", 15000, False, "Kuretase"),
    ("Alkohol", "Alkohol", "BHP", 5000, False, "Kuretase"),
    ("Abbocath", "Abbocath", "BHP", 3500, False, "Kuretase"),
    ("Spuit 10 cc", "Spluit 10 cc", "BHP", 5000, False, "mabok.ngidam"),
    ("Pembalut", "Pembalut", "BHP", 2000, False, "KWITANSI FULL "),
]


def key(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", s.lower())


def clean(s) -> str:
    s = re.sub(r"\s+", " ", str(s).strip())
    return s[:1].upper() + s[1:] if s else s


def num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    return None   # '#N/A', '#VALUE!', teks, dll.


def category_for(row: int, name: str) -> str:
    n = name.lower()
    if n in OVERRIDE:
        return OVERRIDE[n]
    if row <= 119:
        return "Obat"
    if n in JASA:
        return "Jasa"
    if n in TINDAKAN:
        return "Tindakan"
    if n in ALAT_BLOCK:
        return "Alat"
    return "BHP"


def run(src: Path, out_dir: Path) -> dict:
    wb = load_workbook(src, data_only=True)
    ws = wb["HARGA"]
    report = {"source": src.name, "sheets": wb.sheetnames, "skipped": [], "cost_unknown": [],
              "duplicates": [], "extras_missing_in_excel": []}
    products, seen = [], {}
    for r in range(3, ws.max_row + 1):
        raw_name, cost_raw, sale_raw = ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value
        if raw_name is None or not str(raw_name).strip():
            continue
        name = clean(raw_name)
        sale = num(sale_raw)
        if sale is None or sale < 0:
            report["skipped"].append({"row": r, "name": name, "reason": f"harga jual tidak valid: {sale_raw!r}"})
            continue
        sale = int(round(sale))
        k = key(name)
        if k in seen:
            first = seen[k]
            same = first["sale_price"] == sale
            report["duplicates"].append({
                "row": r, "name": name, "price": sale, "kept_row": first["_row"],
                "kept_price": first["sale_price"],
                "note": "harga sama, dilewati" if same else
                        "harga BERBEDA; dipakai baris pertama (sama seperti VLOOKUP Excel)"})
            continue
        cost = num(cost_raw)
        if cost is None or cost <= 0:
            report["cost_unknown"].append({"row": r, "name": name, "cost_raw": None if cost_raw is None else str(cost_raw)})
            cost = None
        else:
            cost = round(cost, 2)
        cat = category_for(r, name)
        p = {"name": name, "category": cat, "unit": None if cat in ("Jasa", "Tindakan", "Lab") else
             ("ampul" if name.lower() in AMPUL else "pcs"),
             "cost_price": cost, "sale_price": sale, "price_editable": False,
             "icon_name": pick_icon(name, cat), "source": f"HARGA!B{r}", "_row": r}
        seen[k] = p
        products.append(p)

    # Item tambahan dari sheet kwitansi (harus benar-benar ada di Excel)
    found = set()
    for sn in wb.sheetnames:
        if sn == "HARGA":
            continue
        for row in wb[sn].iter_rows(min_col=1, max_col=2, values_only=True):
            for v in row:
                if isinstance(v, str):
                    found.add(v.strip().lower())
    for name, src_name, cat, price, editable, sheet in EXTRAS:
        if src_name.lower() not in found:
            report["extras_missing_in_excel"].append(name)
            continue
        if key(name) in seen:
            continue
        products.append({"name": name, "category": cat,
                         "unit": None if cat in ("Jasa", "Tindakan", "Lab") else "pcs",
                         "cost_price": None, "sale_price": price, "price_editable": editable,
                         "icon_name": pick_icon(name, cat), "source": f"kwitansi:{sheet}", "_row": 0})
    for p in products:
        p.pop("_row", None)
    by_cat = {}
    for p in products:
        by_cat[p["category"]] = by_cat.get(p["category"], 0) + 1
    report["counts"] = {"products": len(products), "by_category": by_cat,
                        "from_harga": sum(1 for p in products if p["source"].startswith("HARGA")),
                        "from_kwitansi": sum(1 for p in products if p["source"].startswith("kwitansi")),
                        "price_editable": sum(1 for p in products if p["price_editable"])}
    out = {"generated_from": src.name,
           "categories": [{"name": n, "receipt_label": l, "receipt_order": o, "sort_order": s,
                           "icon_name": CATEGORY_ICON[n]} for n, l, o, s in CATEGORIES],
           "products": products}
    (out_dir / "products.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    (out_dir / "import_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=1), encoding="utf-8")
    return report


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SRC
    report = run(src, ROOT / "seed")
    print(json.dumps(report["counts"], ensure_ascii=False))
    print("duplikat:", [(d["name"], d["note"]) for d in report["duplicates"]])
    print("harga modal tidak diketahui:", len(report["cost_unknown"]), "item")
    print("tidak ditemukan di Excel:", report["extras_missing_in_excel"], "dilewati:", report["skipped"])


if __name__ == "__main__":
    main()
