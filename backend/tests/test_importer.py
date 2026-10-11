"""Importer Excel harus aman terhadap data kosong/#N/A/invalid."""
import json
import sys
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "tools"))
import extract_excel  # noqa: E402
from app.icons import ICON_NAMES  # noqa: E402


def test_importer_survives_broken_workbook(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "HARGA"
    ws["B2"], ws["C2"], ws["D2"] = "OBAT RESEP", "HARGA", "HARGA"
    rows = [("Obat A", "#N/A", 1000), ("Obat B", 0, 2000), ("obat b", 500, 2500), (None, None, None),
            ("Rusak", 100, "#N/A"), ("Teks", "abc", "xyz"), ("Obat C", None, 3000.4), ("  Spasi   Ganda ", 10, 500)]
    for i, (b, c, d) in enumerate(rows, start=3):
        ws.cell(i, 2, b); ws.cell(i, 3, c); ws.cell(i, 4, d)
    src = tmp_path / "rusak.xlsx"
    wb.save(src)
    report = extract_excel.run(src, tmp_path)
    data = json.loads((tmp_path / "products.json").read_text(encoding="utf-8"))
    names = {p["name"]: p for p in data["products"]}
    assert "Obat A" in names and names["Obat A"]["cost_price"] is None          # #N/A -> null, bukan harga
    assert names["Obat B"]["sale_price"] == 2000                                  # duplikat: baris pertama
    assert "Spasi Ganda" in names and names["Obat C"]["sale_price"] == 3000
    assert {s["name"] for s in report["skipped"]} == {"Rusak", "Teks"}
    assert report["duplicates"] and all(isinstance(p["sale_price"], int) for p in data["products"])
    assert not any("N/A" in json.dumps(p) for p in data["products"])


def test_real_seed_is_clean():
    data = json.loads((ROOT / "seed" / "products.json").read_text(encoding="utf-8"))
    prods = data["products"]
    assert len(prods) >= 190
    assert all(isinstance(p["sale_price"], int) and p["sale_price"] >= 0 for p in prods)
    assert all(p["cost_price"] is None or p["cost_price"] > 0 for p in prods)
    assert all(p["icon_name"] in ICON_NAMES for p in prods)
    assert not any("#" in p["name"] for p in prods)
    cats = {c["name"] for c in data["categories"]}
    assert {"Obat", "Jasa", "Tindakan", "Alat"} <= cats and {p["category"] for p in prods} <= cats
    by = {p["name"]: p for p in prods}
    assert by["Konsultasi Dokter"]["sale_price"] == 120000 and by["USG 2D"]["sale_price"] == 80000
    assert by["Kalnex"]["sale_price"] == 7500 and by["Prove D3-5000"]["sale_price"] == 7500
    assert by["Ceftriaxone sodium 1 G Inj"]["unit"] == "ampul" and by["Konsultasi Dokter"]["unit"] is None
