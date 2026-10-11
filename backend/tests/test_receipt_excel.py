"""Bandingkan STRUK APLIKASI vs KWITANSI EXCEL (isi/teks/urutan/format angka)."""
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.receipt import format_amount, format_date_id, number_line, split_title
from app.terbilang import terbilang_rupiah

XLSX = Path(__file__).resolve().parent.parent / "seed" / "data-Alqoyyum-royal-clinic.xlsx"
SAMPLES = ["KWITANSI FULL ", "KWITANSI jasa", "KWITANSI obat", "KWITANSI FULL  (3)", "KWITANSI FULL  (2)"]


def parse_sheet(ws):
    """Ambil baris item, total, terbilang, nama pasien dari kwitansi Excel (nilai cache)."""
    start = end = words_row = patient_row = None
    for r in range(1, 40):
        a = str(ws.cell(r, 1).value or "").strip()
        if a == "Untuk Pembayaran" and start is None:
            start = r
        elif a.startswith("Total") and start and end is None:
            end = r
        elif a == "Terbilang" and words_row is None:
            words_row = r
        elif a == "Nama Pasien" and patient_row is None:
            patient_row = r
    lines = []
    for r in range(start + 1, end):
        name = ws.cell(r, 2).value
        if not name or not str(name).strip():
            continue
        qty, unit = ws.cell(r, 5).value, ws.cell(r, 6).value
        lines.append((str(name).strip().lower(),
                      int(qty) if isinstance(qty, (int, float)) else None,
                      str(unit).strip().lower() if unit else None,
                      format_amount(int(ws.cell(r, 7).value))))
    return {"lines": lines, "total": int(ws.cell(end, 7).value),
            "words": re.sub(r"\s+", " ", str(ws.cell(words_row, 4).value)).strip(),
            "patient": str(ws.cell(patient_row, 4).value).strip()}


@pytest.fixture(scope="module")
def wbv():
    return load_workbook(XLSX, data_only=True)


@pytest.mark.parametrize("sheet", SAMPLES)
def test_app_receipt_matches_excel_receipt(client, kasir_h, pid, wbv, sheet):
    exp = parse_sheet(wbv[sheet])
    items = []
    for name, qty, unit, _ in exp["lines"]:
        items.append({"product_id": pid(name), "quantity": qty or 1})
    r = client.post("/api/transactions", headers=kasir_h,
                    json={"patient_name": exp["patient"], "payment": exp["total"] + 1000, "items": items})
    assert r.status_code == 201, (sheet, r.text)
    tx = r.json()
    assert tx["total"] == exp["total"], f"total beda di {sheet}"
    doc = client.get(f"/api/transactions/{tx['id']}/receipt", headers=kasir_h).json()
    got = sorted((l["name"].lower(), l["qty"], (l["unit"] or "").lower() or None, l["amount_text"])
                 for g in doc["groups"] for l in g["lines"])
    assert got == sorted(exp["lines"]), f"baris struk beda di {sheet}"
    assert doc["total_text"] == format_amount(exp["total"])
    assert doc["words"] == exp["words"], f"terbilang beda di {sheet}"
    assert doc["patient"]["title"] == "Ny." and doc["patient"]["name"] == exp["patient"]


def test_static_receipt_texts_follow_excel(client, kasir_h, pid):
    tx = client.post("/api/transactions", headers=kasir_h,
                     json={"patient_name": "Yayuk Novita Dewi", "payment": 200000,
                           "items": [{"product_id": pid("Konsultasi Dokter"), "quantity": 1},
                                     {"product_id": pid("USG 2D"), "quantity": 1}]}).json()
    d = client.get(f"/api/transactions/{tx['id']}/receipt", headers=kasir_h).json()
    assert d["clinic"]["name"] == "AL-QOYYUM ROYAL CLINIC"
    assert d["clinic"]["address"] == "Jl. Buana Indah No.61 Cihideung-Sadang-Purwakarta"
    assert d["clinic"]["phone_line"] == "Telp : (0264)8464554 / 0813.1314.8633 (WA)"
    assert d["clinic"]["sipd_line"] == "No. SIPD : 446/SIPD-1785/DPMPTSP/2022"
    assert d["title"] == "BUKTI PEMBAYARAN"
    assert d["number_line"] == f"NO : {tx['transaction_number']}"
    assert (d["received_from"]["label"], d["patient"]["label"], d["doctor"]["label"]) == \
        ("Diterima dari", "Nama Pasien", "Nama Dokter")
    assert (d["doctor"]["title"], d["doctor"]["name"]) == ("dr.", "Maulia Mardini, SpOG, MARS")
    assert d["for_payment_label"] == "Untuk Pembayaran" and d["total_label"] == "Total Pembayaran"
    assert (d["payment_method_label"], d["payment_method"], d["words_label"]) == ("Cara Pembayaran", "Cash", "Terbilang")
    assert d["signature"] == {"label": "Diterima oleh:", "name": "Nunung Nurjanah"}
    assert [g["label"] for g in d["groups"]] == ["Jasa"] and len(d["groups"][0]["lines"]) == 2
    assert d["groups"][0]["lines"][0] == {"name": "Konsultasi Dokter", "qty": None, "unit": None,
                                          "amount": 120000, "amount_text": "120.000,00"}
    assert d["words"] == "Duaratus Ribu Rupiah" and d["currency"] == "Rp" and d["paper"]["width_mm"] == 176


def test_groups_follow_excel_labels_and_order(client, kasir_h, pid):
    items = [("Cefixime", 6), ("Nova T IUD", 1), ("Konsultasi Dokter", 1), ("Glukosa Sewaktu", 1),
             ("Kasa", 4), ("Insersi IUD", 1)]
    tx = client.post("/api/transactions", headers=kasir_h, json={"patient_name": "Tn. Samijo", "payment": 10 ** 7,
         "items": [{"product_id": pid(n), "quantity": q} for n, q in items]}).json()
    d = client.get(f"/api/transactions/{tx['id']}/receipt", headers=kasir_h).json()
    assert [g["label"] for g in d["groups"]] == ["Jasa", "Lab", "Alat", "Obat", "BHP"]
    assert [l["name"] for l in d["groups"][0]["lines"]] == ["Konsultasi Dokter", "Insersi IUD"]  # Tindakan tampil di bawah "Jasa"
    assert d["patient"]["title"] == "Tn." and d["patient"]["name"] == "Samijo"
    cef = d["groups"][3]["lines"][0]
    assert (cef["qty"], cef["unit"], cef["amount_text"]) == (6, "pcs", "42.000,00")


def test_formats():
    from datetime import datetime
    assert format_amount(739000) == "739.000,00" and format_amount(1585000) == "1.585.000,00"
    assert format_date_id(datetime(2026, 10, 9)) == "09 Oktober 2026"
    assert format_date_id(datetime(2023, 4, 16)) == "16 April 2023"
    assert number_line("367/1026") == "NO : 367/1026"
    assert split_title("Drg. Aghnia Khairunnisa A") == ("Drg.", "Aghnia Khairunnisa A")
    assert terbilang_rupiah(5500000) == "Lima Juta Limaratus Ribu Rupiah"   # gaya add-in Excel
